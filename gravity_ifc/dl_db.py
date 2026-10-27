"""Load the dead load weights spreadsheet into a flat list of rows.

Spreadsheet layout (sheet "Dead Loads"): header on row 3. Columns are found by
header name: dimension (in) | material | [orientation] | use/component |
additional info | psf | source. Kim's "Combined" sheet added orientation
(Horizontal / Vertical / Both); older sheets have no such column.

A psf cell may be:
  - a number               -> fixed psf for that row
  - a formula like =B40*2  -> psf per inch of the row's dimension (the row is
                              thickness-driven: psf = factor * t)
  - text like "1.5–2.5"    -> a range; psf is None, the text is kept
"""
from __future__ import annotations

import re
from dataclasses import dataclass, asdict
from pathlib import Path

import openpyxl

SHEET = "Dead Loads"
HEADER_ROW = 3


@dataclass
class DbRow:
    id: str                 # "R40" = spreadsheet row 40, stable while the sheet is
    t_in: float | None      # dimension in inches, None for "--"
    material: str
    use: str
    info: str
    psf: float | None       # psf at t_in (None for ranges)
    psf_per_in: float | None  # set when the row scales with thickness
    psf_text: str | None    # original text when psf is a range
    source: str
    orientation: str = ""   # Horizontal / Vertical / Both, "" when the sheet has no column

    def psf_at(self, t_in: float | None) -> float | None:
        """psf for a layer of thickness t_in using this row."""
        if self.psf_per_in is not None and t_in:
            return round(self.psf_per_in * t_in, 2)
        return self.psf

    def label(self) -> str:
        dim = f'{self.t_in:g}" ' if self.t_in is not None else ""
        info = f" ({self.info})" if self.info and self.info != "--" else ""
        return f"{dim}{self.material} - {self.use}{info}"

    def to_dict(self) -> dict:
        d = asdict(self)
        d["label"] = self.label()
        return d


_CELL_REF = re.compile(r"\$?[A-Z]{1,2}\$?\d+")


def _per_inch_factor(formula: str) -> float | None:
    """'=B40*0.4*8' -> 3.2. Formulas are linear in the dimension cell."""
    expr = formula.lstrip("=")
    if not _CELL_REF.search(expr):
        return None
    expr = _CELL_REF.sub("1", expr)
    if not re.fullmatch(r"[\d.\s*/+\-()]+", expr):
        return None
    return float(eval(expr, {"__builtins__": {}}, {}))  # digits and operators only


# header text (lower-cased, startswith) -> field; first match wins
_HEADERS = {"dimension": "dim", "material": "mat", "orientation": "orient", "use": "use",
            "additional": "info", "psf": "psf", "source": "src"}


def _columns(ws) -> dict[str, int]:
    cols: dict[str, int] = {}
    for c in range(1, ws.max_column + 1):
        h = str(ws.cell(HEADER_ROW, c).value or "").strip().lower()
        for prefix, field in _HEADERS.items():
            if h.startswith(prefix) and field not in cols:
                cols[field] = c
    if not cols:  # no header row: the original fixed layout, B..G
        cols = dict(zip(("dim", "mat", "use", "info", "psf", "src"), range(2, 8)))
    missing = {"mat", "psf"} - cols.keys()
    if missing:
        raise ValueError(f"dead load sheet: header row {HEADER_ROW} has no {sorted(missing)} column")
    return cols


def load_db(path: str | Path) -> list[DbRow]:
    wb = openpyxl.load_workbook(path, data_only=False)
    ws = wb[SHEET] if SHEET in wb.sheetnames else wb.worksheets[0]
    cols = _columns(ws)
    rows: list[DbRow] = []
    for r in range(HEADER_ROW + 1, ws.max_row + 1):
        cell = lambda f: ws.cell(r, cols[f]).value if f in cols else None  # noqa: E731
        dim, mat, orient, use, info, psf_cell, src = (cell(f) for f in ("dim", "mat", "orient", "use", "info", "psf", "src"))
        if not mat or not isinstance(mat, str) or psf_cell is None:
            continue  # notes / blank rows
        t_in = float(dim) if isinstance(dim, (int, float)) else None
        psf = per_in = text = None
        if isinstance(psf_cell, (int, float)):
            psf = float(psf_cell)
        elif isinstance(psf_cell, str) and psf_cell.startswith("="):
            per_in = _per_inch_factor(psf_cell)
            psf = round(per_in * t_in, 2) if per_in is not None and t_in else None
        else:
            text = str(psf_cell)
        rows.append(DbRow(f"R{r}", t_in, mat.strip(), str(use or "").strip(), str(info or "").strip(),
                          psf, per_in, text, str(src or "").strip(), str(orient or "").strip()))
    return rows


# Rows the matcher needs that the spreadsheet doesn't have yet (zero-weight
# layers, per the 10/27 decision that these should be rows). Proposed for Kim's
# database; kept here so matching works today.
SUPPLEMENTAL = [
    DbRow("X1", None, "air space", "cavity", "zero weight", 0.0, None, None, "Proposed"),
    DbRow("X2", None, "vapor/air barrier", "membrane", "zero weight", 0.0, None, None, "Proposed"),
    DbRow("X3", None, "paint / coating", "finish", "zero weight", 0.0, None, None, "Proposed"),
    DbRow("X4", None, "load allowance", "allowance", "MEP", None, None, None, "Proposed"),
    DbRow("X5", None, "load allowance", "allowance", "concrete ponding", None, None, None, "Proposed"),
    DbRow("X6", None, "load allowance", "allowance", "future finishes", None, None, None, "Proposed"),
]


if __name__ == "__main__":
    import sys
    db = load_db(sys.argv[1])
    print(len(db), "rows")
    for row in db:
        if row.psf_per_in or row.psf_text:
            print(row.id, row.label(), row.psf, row.psf_per_in, row.psf_text)
