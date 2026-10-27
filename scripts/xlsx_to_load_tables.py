"""Convert the two load spreadsheets into the LoadTable JSON the Loads page reads.

Usage: python scripts/xlsx_to_load_tables.py <live.xlsx> <dead.xlsx>
Writes frontend/src/ltl/data/live-loads.json and dead-loads.json. The shape matches `LoadTable`
in frontend/src/ltl/loads.ts (which parses the same sheets in the browser on upload).
"""
import json
import os
import sys
from pathlib import Path

import openpyxl

OUT = Path(__file__).resolve().parent.parent / "frontend" / "src" / "ltl" / "data"


def clean(v):
    if v is None:
        return None
    if isinstance(v, bool):
        return v
    if isinstance(v, (int, float)):
        return int(v) if float(v).is_integer() else v
    s = str(v).strip()
    return s or None


def convert(path, sheet_re, header_needles, notes_sheet=True):
    wb = openpyxl.load_workbook(path, data_only=True)  # cached results, not the formula text
    main = next(ws for ws in wb if sheet_re in ws.title.lower())
    rows = [[clean(c) for c in r] for r in main.iter_rows(values_only=True)]
    h = next(i for i, r in enumerate(rows)
             if all(any(isinstance(c, str) and n in c.lower() for c in r) for n in header_needles))
    header = rows[h]
    first = next(i for i, c in enumerate(header) if c is not None)
    last = max(i for i, c in enumerate(header) if c is not None)
    columns = [str(c) if c is not None else f"Column {first + i + 1}" for i, c in enumerate(header[first:last + 1])]
    above = [" ".join(str(c) for c in r if c is not None) for r in rows[:h] if any(c is not None for c in r)]
    comments = {}
    for row in main.iter_rows():
        for c in row:
            if c.comment and c.comment.text.strip():
                comments.setdefault(c.row, []).append(c.comment.text.strip())
    data, notes_per_row = [], []
    i = h + 1
    while i < len(rows) and any(c is not None for c in rows[i]):
        data.append([rows[i][first + k] if first + k < len(rows[i]) else None for k in range(len(columns))])
        notes_per_row.append(comments.get(i + 1))
        i += 1
    footer = [" ".join(str(c) for c in r if c is not None) for r in rows[i:] if any(c is not None for c in r)]
    notes = []
    other = next((ws for ws in wb if ws is not main and "note" in ws.title.lower()), None)
    if other:
        for n, r in enumerate(other.iter_rows(values_only=True)):
            r = [clean(c) for c in r]
            if (r[0] is None and r[1] is None) or (n == 0 and str(r[0]).lower() == "reference"):
                continue
            notes.append({"ref": str(r[0] or ""), "text": str(r[1] or "")})
    return {
        "fileName": os.path.basename(path), "sheetName": main.title,
        "title": above[0] if above else main.title, "captions": above[1:],
        "columns": columns, "rows": data, "rowNotes": notes_per_row, "footer": footer, "notes": notes,
    }


if __name__ == "__main__":
    live, dead = sys.argv[1:3]
    OUT.mkdir(parents=True, exist_ok=True)
    for name, table in (("live-loads", convert(live, "live load", ["occupancy", "psf"])),
                        ("dead-loads", convert(dead, "dead load", ["material", "psf"]))):
        (OUT / f"{name}.json").write_text(json.dumps(table, ensure_ascii=False, indent=1) + "\n")
        flagged = sum(1 for n in table["rowNotes"] if n)
        print(f"{name}: {len(table['rows'])} rows, {flagged} flagged, {len(table['notes'])} notes, {len(table['footer'])} footer lines")
