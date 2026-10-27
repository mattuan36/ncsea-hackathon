"""Materials table: Kim's spreadsheet import, edits, user rows (docker Postgres)."""
from pathlib import Path

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app import materials
from app.main import app
from dl_db import load_db


def _sheet(path: Path, with_orientation: bool) -> Path:
    """A two-row dead load sheet in either layout (with or without Kim's orientation column)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Dead Loads"
    head = ["dimension (in)", "material", "orientation", "use/component", "additional info", "psf", "source"]
    rows = [[8, "cmu", "Vertical", "wall", "solid grout", 75, "ASCE C3.1-1a"],
            [1, "Tectum", "Horizontal", "roof deck", "2 psf per in.", "=B5*2", "Firm DL"]]
    if not with_orientation:
        head.pop(2)
        rows = [r[:2] + r[3:] for r in rows]
    for c, h in enumerate(head, start=2):
        ws.cell(3, c, h)
    for r, row in enumerate(rows, start=4):
        for c, v in enumerate(row, start=2):
            ws.cell(r, c, v)
    wb.save(path)
    return path


@pytest.mark.parametrize("with_orientation", [True, False])
def test_loader_reads_columns_by_header(tmp_path, with_orientation):
    cmu, tectum = load_db(_sheet(tmp_path / "dl.xlsx", with_orientation))
    assert (cmu.material, cmu.use, cmu.info, cmu.psf, cmu.source) == ("cmu", "wall", "solid grout", 75, "ASCE C3.1-1a")
    assert cmu.orientation == ("Vertical" if with_orientation else "")
    assert (tectum.psf_per_in, tectum.psf) == (2.0, 2.0)


def test_materials_crud_and_reimport_keeps_edits(tmp_path, monkeypatch):
    monkeypatch.setenv("DEAD_LOAD_XLSX", str(_sheet(tmp_path / "dl.xlsx", True)))
    with TestClient(app) as client:
        assert client.post("/api/materials/import").status_code == 200
        rows = {m["id"]: m for m in client.get("/api/materials").json()}
        assert rows["R4"]["material"] == "cmu" and rows["R4"]["psf"] == 75
        assert "X1" in rows  # proposed zero-weight rows ride along

        r = client.patch("/api/materials/R4", json={"psf": 80})
        assert r.json()["psf"] == 80 and r.json()["edited"]
        added = client.post("/api/materials", json={"material": "CLT panel", "t_in": 5.5, "psf_per_in": 2.9, "use": "floor"}).json()
        assert added["custom"] and added["id"].startswith("U") and added["psf"] == 15.95

        client.post("/api/materials/import")
        rows = {m["id"]: m for m in client.get("/api/materials").json()}
        assert rows["R4"]["psf"] == 80, "re-import overwrote a user edit"
        assert added["id"] in rows

        assert client.delete(f"/api/materials/{added['id']}").status_code == 204
        # put R4 back as the spreadsheet has it so later runs start clean
        client.patch("/api/materials/R4", json={"psf": 75})


def test_matcher_rows_skip_supplemental(monkeypatch):
    with TestClient(app):
        ids = {r.id for r in materials.matcher_rows()}
    assert not ids & materials.SUPPLEMENTAL_IDS
