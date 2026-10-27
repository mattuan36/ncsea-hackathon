"""Read Part 1 hazards from a Hazard_Exchange XLSX or a Revit-import CSV.

The XLSX contract is BP-HZ-1. The CSV contract is the Lateral page export:
parameter_name, value, units, description, preliminary, source_status, engineer_confirmed.
Blank hazard cells stay missing. They are never read as zero.
"""
import csv
from io import BytesIO, StringIO
from openpyxl import load_workbook

REQUIRED = {'schema_version', 'asce_edition', 'source_status', 'engineer_confirmed',
            's1_g', 'sds_g', 'sd1_g', 'tl_sec'}
NUMERIC = {'s1_g', 'sds_g', 'sd1_g', 'tl_sec'}
RISK_CATEGORIES = ('I', 'II', 'III', 'IV')
SOURCE_STATUSES = ('SYNTHETIC_DEMO', 'ASCE_API_UNVERIFIED', 'SKYCIV_GETSITEDATA')
REVIT_MAP = {
    'BP_HZ_PROJECT_NUMBER': 'project_number',
    'BP_HZ_ASCE_EDITION': 'asce_edition',
    'BP_HZ_RISK_CATEGORY': 'risk_category',
    'BP_HZ_SITE_CLASS': 'site_class',
    'BP_HZ_SDC': 'sdc',
    'BP_HZ_SDS': 'sds_g',
    'BP_HZ_SD1': 'sd1_g',
    'BP_HZ_S1': 's1_g',
    'BP_HZ_TL': 'tl_sec',
    'BP_HZ_W_PRELIM': 'weight_kips',
}


def _text(v):
    return '' if v is None else str(v).strip()


def _edition(value):
    compact = _text(value).lower().replace(' ', '')
    if compact in ('7-22', 'asce7-22'):
        return '7-22'
    raise ValueError(f"Part 3 supports ASCE 7-22 only; workbook is {value!r}. "
                     'Regenerate the Part 1 file for ASCE 7-22.')


def _number(value, key):
    if isinstance(value, str):
        value = value.strip()
        if value == '':
            value = None
    if value is None or isinstance(value, bool) or not isinstance(value, (int, float, str)):
        raise ValueError(f'Required {key} must be a numeric hazard value; missing is not zero.')
    if isinstance(value, str):
        try:
            value = float(value)
        except ValueError:
            raise ValueError(f'Required {key} must be a numeric hazard value; missing is not zero.') from None
    if value < 0:
        raise ValueError(f'Required {key} must be a numeric hazard value; missing is not zero.')
    return float(value)


def _finish(values, rows):
    edition = _edition(values.get('asce_edition'))
    numbers = {k: _number(values.get(k), k) for k in NUMERIC}
    if numbers['tl_sec'] <= 0:
        raise ValueError('TL must be positive.')
    status = _text(values.get('source_status'))
    if status not in SOURCE_STATUSES:
        raise ValueError('Unrecognized data source status; no silent source approval.')
    warnings = ['Unverified or synthetic hazard data; engineer must independently review.']
    rc = _text(values.get('risk_category')).upper()
    if rc not in RISK_CATEGORIES:
        warnings.append('Risk Category missing or not I-IV in Part 1 file; engineer must enter it.')
        rc = ''
    weight = values.get('weight_kips', None)
    if weight is None or _text(weight) == '':
        weight = None
    else:
        weight = _number(weight, 'weight_kips')
        if weight <= 0:
            raise ValueError('W must be positive when provided; missing is not zero.')
    return {'project_number': _text(values.get('project_number')),
            'asce_edition': edition, 'source_status': status,
            'engineer_confirmed': _text(values.get('engineer_confirmed')).upper() == 'YES',
            'sds': numbers['sds_g'], 'sd1': numbers['sd1_g'],
            's1': numbers['s1_g'], 'tl': numbers['tl_sec'],
            'risk_category': rc, 'site_class': _text(values.get('site_class')),
            'part1_sdc': _text(values.get('sdc')).upper(),
            'weight_kips': weight,
            'exchange_rows': rows, 'warnings': warnings, 'warning': warnings[0]}


def _read_xlsx(raw: bytes) -> dict:
    try:
        wb = load_workbook(BytesIO(raw), read_only=True, data_only=True)
    except Exception:
        raise ValueError('Invalid or unreadable Excel workbook.') from None
    if 'Hazard_Exchange' not in wb.sheetnames:
        raise ValueError('Required sheet Hazard_Exchange not found. Upload an XLSX from Part 1.')
    ws = wb['Hazard_Exchange']
    values = {}
    rows = []
    for row in ws.iter_rows(min_row=2, max_col=5, values_only=True):
        key, val, unit, source, review = (list(row) + [None] * 5)[:5]
        if key is None:
            continue
        if not isinstance(key, str) or len(key) > 80:
            raise ValueError('Invalid exchange key.')
        if key in values:
            raise ValueError('Duplicate exchange key: ' + key)
        values[key] = val
        rows.append([key, val, unit, source, review])
    missing = REQUIRED - set(values)
    if missing:
        raise ValueError('Required keys missing: ' + ', '.join(sorted(missing)))
    if values['schema_version'] != 'BP-HZ-1':
        raise ValueError('Unsupported exchange schema version.')
    return _finish(values, rows)


def _cell(row, header, name):
    key = header.get(name)
    return '' if key is None else _text(row.get(key))


def _read_revit_csv(table, header) -> dict:
    values = {}
    rows = []
    sds_source = ''
    sds_confirmed = 'NO'
    for row in table:
        name = _cell(row, header, 'parameter_name')
        if not name:
            continue
        if len(name) > 80:
            raise ValueError('Invalid exchange key.')
        if name in {r[0] for r in rows}:
            raise ValueError('Duplicate exchange key: ' + name)
        value = _cell(row, header, 'value')
        unit = _cell(row, header, 'units')
        source = _cell(row, header, 'source_status')
        confirmed = _cell(row, header, 'engineer_confirmed') or 'NO'
        rows.append([name, value, unit, source, confirmed])
        mapped = REVIT_MAP.get(name)
        if mapped:
            values[mapped] = value
        if name == 'BP_HZ_SDS':
            sds_source = source
            sds_confirmed = confirmed
    values['source_status'] = sds_source
    values['engineer_confirmed'] = sds_confirmed
    missing = (REQUIRED - {'schema_version'}) - set(values)
    if missing:
        raise ValueError('Required keys missing: ' + ', '.join(sorted(missing)))
    parsed = _finish(values, rows)
    review = 'YES' if parsed['engineer_confirmed'] else 'NO'
    parsed['exchange_rows'] = rows + [
        ['sds_g', parsed['sds'], 'g', parsed['source_status'], review],
        ['sd1_g', parsed['sd1'], 'g', parsed['source_status'], review],
        ['s1_g', parsed['s1'], 'g', parsed['source_status'], review],
        ['tl_sec', parsed['tl'], 'sec', parsed['source_status'], review],
    ]
    return parsed


def _read_csv(raw: bytes) -> dict:
    try:
        text = raw.decode('utf-8-sig')
    except UnicodeDecodeError:
        raise ValueError('CSV is not UTF-8 text.') from None
    reader = csv.DictReader(StringIO(text))
    if not reader.fieldnames:
        raise ValueError('CSV has no header row. Expected parameter_name.')
    header = {_text(name).lower(): name for name in reader.fieldnames if name}
    if 'parameter_name' not in header:
        raise ValueError('Unrecognized CSV header. Expected parameter_name.')
    return _read_revit_csv(reader, header)


def read_exchange(raw: bytes) -> dict:
    if len(raw) > 6_000_000:
        raise ValueError('File exceeds 6MB input size limit.')
    if raw[:2] == b'PK':
        return _read_xlsx(raw)
    return _read_csv(raw)
