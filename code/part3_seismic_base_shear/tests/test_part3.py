"""Part 3 tests. Run: python -m pytest -q (from part3_seismic_base_shear)."""
from io import BytesIO
import base64
import pytest
from openpyxl import Workbook, load_workbook
import code_tables as ct
from calculation import SeismicInput, calculate
from workbook import build_calculation_xlsx
from exchange import read_exchange

BASE = {'project_number': 'DEMO-1', 'source_status': 'SYNTHETIC_DEMO', 'sds': .85, 'sd1': .55, 's1': .45,
        'tl': 8., 'risk_category': 'II', 'site_class': 'D', 'part1_sdc': 'D', 'hn_ft': 60., 'weight_kips': 1500.}


def single(sid, d, **kw): return {'direction': d, 'combination': 'single', 'system_ids': [sid], **kw}


def inputs(x=None, y=None, **change):
    raw = dict(BASE, directions=[x or single('B2', 'X'), y or single('B2', 'Y')]); raw.update(change)
    return SeismicInput.model_validate(raw)


def has(warnings, text): return any(text in w for w in warnings)


# ---------- tables ----------
def test_table_12_2_1_loaded_from_verified_json():
    t = ct.table_12_2_1()
    assert len(t) == 92
    assert (t['C1']['R'], t['C1']['omega0'], t['C1']['Cd']) == (8, 3, 5.5)
    assert t['B3']['height_limits_ft']['F'] == 'NP'


@pytest.mark.parametrize('rc,ie', [('I', 1.0), ('II', 1.0), ('III', 1.25), ('IV', 1.5)])
def test_ie_from_risk_category(rc, ie):
    assert calculate(inputs(risk_category=rc))['ie'] == ie


@pytest.mark.parametrize('sds,sd1,s1,rc,sdc', [
    (.10, .05, .03, 'II', 'A'), (.167, .05, .03, 'II', 'B'), (.167, .05, .03, 'IV', 'C'),
    (.33, .05, .03, 'II', 'C'), (.33, .05, .03, 'IV', 'D'), (.2, .133, .1, 'II', 'C'),
    (.2, .2, .1, 'I', 'D'), (1.5, 1., .75, 'II', 'E'), (1.5, 1., .75, 'IV', 'F')])
def test_sdc(sds, sd1, s1, rc, sdc):
    assert ct.seismic_design_category(sds, sd1, s1, rc)['sdc'] == sdc


@pytest.mark.parametrize('sd1,cu', [(.5, 1.4), (.3, 1.4), (.25, 1.45), (.2, 1.5), (.175, 1.55), (.1, 1.7), (.05, 1.7)])
def test_cu(sd1, cu): assert ct.cu(sd1) == pytest.approx(cu)


# ---------- period ----------
def test_period_capped_at_cu_ta():
    r = calculate(inputs(x=single('C1', 'X', period_analysis_sec=1.2, moment_frames_resist_100pct=True)))['directions']['X']
    ta = 0.028 * 60 ** 0.8
    assert r['Ta'] == pytest.approx(ta) and r['T_used'] == pytest.approx(1.4 * ta)
    assert r['cs'] == pytest.approx(.55 / (1.4 * ta * 8))


def test_period_analysis_used_when_below_cap_and_ta_when_missing():
    r = calculate(inputs(x=single('B2', 'X', period_analysis_sec=0.3)))['directions']
    assert r['X']['T_used'] == pytest.approx(0.3)
    assert r['Y']['T_used'] == pytest.approx(0.02 * 60 ** 0.75)


def test_moment_frame_ct_requires_confirmation():
    res = calculate(inputs(x=single('C1', 'X')))
    assert res['directions']['X']['period_type'] == 'all_other' and has(res['warnings'], '100%')
    res = calculate(inputs(x=single('C1', 'X', moment_frames_resist_100pct=True)))
    assert res['directions']['X']['period_type'] == 'steel_mf'


# ---------- combinations ----------
def test_horizontal_combination_least_r():
    x = {'direction': 'X', 'combination': 'horizontal', 'system_ids': ['B2', 'B26']}  # R 6 vs 8
    d = calculate(inputs(x=x))['directions']['X']
    assert (d['R'], d['omega0'], d['Cd']) == (6, 2, 5)


def test_vertical_lower_system_has_lower_r():
    y = {'direction': 'Y', 'combination': 'vertical', 'system_ids': ['B2', 'A1']}  # upper R6, lower R5
    d = calculate(inputs(y=y))['directions']['Y']
    assert d['R'] == 5 and d['transfer_amplification'] == pytest.approx(1.2)
    assert d['cs'] == pytest.approx(.85 / 5) and d['cs_upper_system'] == pytest.approx(.85 / 6)


def test_vertical_upper_system_has_lower_r():
    y = {'direction': 'Y', 'combination': 'vertical', 'system_ids': ['A16', 'C5']}  # upper R6.5, lower R8
    d = calculate(inputs(y=y))['directions']['Y']
    assert d['R'] == 6.5 and d['omega0'] == 3 and d['transfer_amplification'] == 1.0


@pytest.mark.parametrize('bad', [
    {'direction': 'X', 'combination': 'single', 'system_ids': ['B2', 'C1']},
    {'direction': 'X', 'combination': 'vertical', 'system_ids': ['B2']},
    {'direction': 'X', 'combination': 'single', 'system_ids': ['Z9']}])
def test_invalid_direction_inputs_rejected(bad):
    with pytest.raises(Exception): inputs(x=bad)


def test_requires_both_directions():
    with pytest.raises(Exception): inputs(y=single('B2', 'X'))


# ---------- limits & warnings (never block) ----------
def test_not_permitted_and_height_warn_but_calculate():
    res = calculate(inputs(hn_ft=200, x=single('C4', 'X'), y=single('B2', 'Y')))
    assert has(res['warnings'], 'NOT PERMITTED') and has(res['warnings'], 'exceeds the 160 ft')
    assert res['directions']['X']['cs'] > 0


def test_footnote_q_min_height():
    assert has(calculate(inputs(hn_ft=50, x=single('B5', 'X')))['warnings'], 'footnote q')


def test_sdc_mismatch_warning():
    assert has(calculate(inputs(part1_sdc='C'))['warnings'], 'SDC MISMATCH')


# ---------- Cs equations ----------
def test_long_period_and_high_s1_minimum():
    d = calculate(inputs(sd1=.6, s1=.8, tl=4, hn_ft=900,
                         x=single('C1', 'X', period_analysis_sec=12, moment_frames_resist_100pct=True)))['directions']['X']
    assert d['T_used'] > 4 and d['cs_upper'] == pytest.approx(.6 * 4 / (d['T_used'] ** 2 * 8))
    assert d['cs_min_high_s1'] == pytest.approx(.5 * .8 / 8) and d['governing'] == 'Minimum Cs'


def test_sds_cap_applied_only_when_all_conditions_met():
    ok = dict(sds=1.6, sd1=1.1, s1=.7, hn_ft=40, stories_above_base=3, site_class='C', part1_sdc='',
              apply_sds_cap_12_8_1_3=True, confirm_regular_no_irregularities=True, confirm_rho_equals_1=True)
    res = calculate(inputs(x=single('B23', 'X'), **ok))
    assert res['directions']['X']['sds_used'] == ct.SDS_CAP_12_8_1_3
    if not ct.SDS_CAP_VERIFIED: assert has(res['warnings'], 'NOT VERIFIED')
    for change in ({'risk_category': 'III'}, {'stories_above_base': 6}, {'site_class': 'E'},
                   {'confirm_rho_equals_1': False}):
        assert calculate(inputs(x=single('B23', 'X'), **dict(ok, **change)))['directions']['X']['sds_used'] == 1.6


def test_no_weight_means_symbolic_only():
    d = calculate(inputs(weight_kips=None))['directions']['X']
    assert d['base_shear_kips'] is None and d['symbolic_base_shear'].startswith('V = ')


# ---------- exchange ----------
def mock_exchange(edition='7-22', **extra):
    values = {'schema_version': 'BP-HZ-1', 'project_number': 'DEMO-1', 'asce_edition': edition,
              'source_status': 'SYNTHETIC_DEMO', 'engineer_confirmed': 'NO', 'risk_category': 'III',
              'site_class': 'D', 'sdc': 'D', 'sds_g': .85, 'sd1_g': .55, 's1_g': .45, 'tl_sec': 8., **extra}
    wb = Workbook(); ws = wb.active; ws.title = 'Hazard_Exchange'
    ws.append(['key', 'value', 'unit', 'source', 'review'])
    for k, v in values.items(): ws.append([k, v, '', 'DEMO', 'NOT FOR DESIGN'])
    out = BytesIO(); wb.save(out); return out.getvalue()


def test_exchange_reads_risk_category_and_rejects_7_16():
    ex = read_exchange(mock_exchange())
    assert ex['risk_category'] == 'III' and ex['part1_sdc'] == 'D' and ex['sds'] == .85
    with pytest.raises(ValueError, match='7-22 only'): read_exchange(mock_exchange('7-16'))


def test_reject_missing_workbook_sheet():
    with pytest.raises(ValueError, match='Hazard_Exchange'):
        wb = Workbook(); f = BytesIO(); wb.save(f); read_exchange(f.getvalue())


SAMPLE_CSV = __import__('pathlib').Path(__file__).resolve().parents[1] / 'samples' / 'Riverside_Office_Building_Hazard_Studio_Revit_Import.csv'


def test_revit_csv_sample_is_the_default_hazard():
    ex = read_exchange(SAMPLE_CSV.read_bytes())
    assert ex['project_number'] == 'Riverside_Office_Building'
    assert ex['source_status'] == 'SKYCIV_GETSITEDATA'
    assert ex['risk_category'] == 'II' and ex['site_class'] == 'D' and ex['part1_sdc'] == 'D'
    assert ex['sds'] == pytest.approx(1.09) and ex['sd1'] == pytest.approx(1.13)
    assert ex['s1'] == pytest.approx(0.6) and ex['tl'] == pytest.approx(12)
    assert ex['weight_kips'] is None and ex['engineer_confirmed'] is False
    assert ex['exchange_rows'][-4][0] == 'sds_g' and ex['exchange_rows'][-4][1] == pytest.approx(1.09)


def test_revit_csv_rejects_other_editions_and_blank_sds():
    text = SAMPLE_CSV.read_text(encoding='utf-8').replace('ASCE 7-22', 'ASCE 7-16')
    with pytest.raises(ValueError, match='7-22 only'):
        read_exchange(text.encode())
    blank = text.replace('ASCE 7-16', 'ASCE 7-22').replace('BP_HZ_SDS,1.090', 'BP_HZ_SDS,')
    with pytest.raises(ValueError, match='missing is not zero'):
        read_exchange(blank.encode())


# ---------- workbook ----------
def test_workbook_formulas_and_cached_results():
    inp = inputs(x=single('C1', 'X', period_analysis_sec=1.2, moment_frames_resist_100pct=True),
                 y={'direction': 'Y', 'combination': 'vertical', 'system_ids': ['B2', 'A1']})
    res = calculate(inp)
    b = build_calculation_xlsx(inp, read_exchange(mock_exchange())['exchange_rows'])
    f = load_workbook(BytesIO(b)); v = load_workbook(BytesIO(b), data_only=True)
    ws_f, ws_v = f['Cs & Base Shear'], v['Cs & Base Shear']
    rows = {ws_f.cell(r, 1).value: r for r in range(1, ws_f.max_row + 1)}
    r_cs, r_v = rows['Cs GOVERNING'], rows['Base shear V = Cs W']
    assert ws_f[f'B{r_cs}'].value.startswith('=MAX(MIN(')
    assert f['Inputs']['B9'].value.startswith("='Imported ASCE'!B")
    for col, d in (('B', 'X'), ('C', 'Y')):
        assert ws_v[f'{col}{r_cs}'].value == pytest.approx(res['directions'][d]['cs'])
        assert ws_v[f'{col}{r_v}'].value == pytest.approx(res['directions'][d]['base_shear_kips'])


# ---------- API ----------
def test_api_endpoints():
    pytest.importorskip('fastapi'); pytest.importorskip('httpx')
    from fastapi.testclient import TestClient
    from app import app
    client = TestClient(app)
    assert len(client.get('/api/systems').json()['systems']) == 92
    body = inputs().model_dump()
    assert client.post('/api/calculate', json=body).status_code == 200
    resp = client.post('/api/export', json={'inputs': body, 'exchange_rows': []})
    assert resp.status_code == 200 and base64.b64decode(resp.json()['xlsx_base64']).startswith(b'PK')
    up = client.post('/api/import', files={'xlsx': ('h.xlsx', mock_exchange(), 'application/octet-stream')})
    assert up.status_code == 200 and up.json()['risk_category'] == 'III'
    csv_up = client.post('/api/import', files={'xlsx': ('h.csv', SAMPLE_CSV.read_bytes(), 'text/csv')})
    assert csv_up.status_code == 200 and csv_up.json()['project_number'] == 'Riverside_Office_Building'
    assert client.get('/api/default-hazard').json()['sds'] == pytest.approx(1.09)
