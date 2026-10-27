"""Regenerate data/table_12_2_1.json from the team-verified ASCE 7-22 table workbook.

Usage: python scripts/build_table_json.py path/to/ASCE7-22_Seismic_Tables.xlsx
Re-run whenever the verified workbook changes; never hand-edit the JSON.
"""
import json, re, sys, hashlib
from pathlib import Path
from openpyxl import load_workbook

ID_RE = re.compile(r'^[A-H]\d*$')
HEIGHT_COLS = ['B', 'C', 'D', 'E', 'F']


def _limit(v):
    if v in ('NL', 'NP'):
        return v
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v)
    raise ValueError(f'Unexpected height-limit value: {v!r}')


def build(src: Path) -> dict:
    raw = src.read_bytes()
    ws = load_workbook(src, read_only=True, data_only=True)['Table 12.2-1']
    systems = []
    for row in ws.iter_rows(min_row=5, values_only=True):
        sid = row[0]
        if not isinstance(sid, str) or not ID_RE.match(sid):
            continue
        r, om, cd = row[5], row[6], row[7]
        for v in (r, om, cd):
            if not isinstance(v, (int, float)) or v <= 0:
                raise ValueError(f'{sid}: R/Omega0/Cd must be positive numbers')
        systems.append({
            'id': sid, 'category': row[1], 'name': row[2],
            'system_footnotes': [s.strip() for s in (row[3] or '').split(',') if s.strip()],
            'detailing_sections': row[4], 'R': float(r), 'omega0': float(om), 'Cd': float(cd),
            'height_limits_ft': dict(zip(HEIGHT_COLS, (_limit(v) for v in row[8:13]))),
            'other_footnotes': row[13] or '',
        })
    if len({s['id'] for s in systems}) != len(systems):
        raise ValueError('Duplicate system IDs')
    return {'standard': 'ASCE/SEI 7-22', 'table': '12.2-1',
            'source_workbook': src.name, 'source_sha256': hashlib.sha256(raw).hexdigest(),
            'verification': 'Reported complete by B&P team; record reviewer/date in the workbook Verification Log.',
            'systems': systems}


if __name__ == '__main__':
    out = Path(__file__).resolve().parents[1] / 'data' / 'table_12_2_1.json'
    data = build(Path(sys.argv[1]))
    out.write_text(json.dumps(data, indent=1, ensure_ascii=False), encoding='utf-8')
    print(f"Wrote {len(data['systems'])} systems to {out}")
